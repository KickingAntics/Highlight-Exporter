"""Layers 2-3: the Word output itself, and the safety rules."""
import getpass
import hashlib
import os
import re
import socket
import subprocess

import pymupdf
import pytest
from docx import Document
from docx.enum.text import WD_COLOR_INDEX

from highlight_export import cli
from highlight_export.extract import PdfProblem, extract
from highlight_export.write_docx import write_combined, write_single

from pdfmaker import GREEN, PARA1, PARA2, PINK, YELLOW, hl, new_doc, put, save


def make(path, phrases=(("written limits on government power", YELLOW),), pages=1, text=PARA1):
    d = new_doc()
    for i in range(pages):
        p = d.new_page()
        put(p, 72, 100, text)
    for ph, color in phrases:
        hl(d[0], ph, color=color)
    return save(d, path)


def settings(**kw):
    s = dict(cli.DEFAULTS)
    s.update(kw)
    return s


def quiet(*a, **k):
    pass


def excerpt_runs(docx_path):
    doc = Document(docx_path)
    out = []
    for par in doc.paragraphs:
        if par.style.name == "Excerpt Text":
            out.append([(r.text, r.font.highlight_color, r.bold) for r in par.runs])
    return doc, out


# ---------------------------------------------------------------- Layer 2
def test_emphasis_only_on_highlighted_text(tmp_path):
    pdf = make(tmp_path / "a.pdf", (("framers of the constitution", GREEN), ("written limits", PINK)))
    res = extract(pdf)
    out = tmp_path / "o.docx"
    write_single(out, res, "a.pdf", "highlight")
    _, paras = excerpt_runs(out)
    marked = [(t, c) for p in paras for t, c, b in p if c is not None]
    assert marked == [("framers of the constitution", WD_COLOR_INDEX.BRIGHT_GREEN),
                      ("written limits", WD_COLOR_INDEX.PINK)]
    assert all(b is None for p in paras for t, c, b in p)  # nothing else bold
    plain = "".join(t for p in paras for t, c, b in p)
    assert plain == " ".join(PARA1)  # context is the PDF text, unchanged


def test_bold_mode_uses_bold_not_highlight(tmp_path):
    res = extract(make(tmp_path / "a.pdf"))
    out = tmp_path / "o.docx"
    write_single(out, res, "a.pdf", "bold")
    _, paras = excerpt_runs(out)
    assert [t for p in paras for t, c, b in p if b] == ["written limits on government power"]
    assert all(c is None for p in paras for t, c, b in p)


def test_every_excerpt_has_pdf_page_and_counts_match(tmp_path):
    res = extract(make(tmp_path / "a.pdf", (("written limits", YELLOW), ("framers", YELLOW))))
    out = tmp_path / "o.docx"
    write_single(out, res, "a.pdf")
    doc = Document(out)
    heads = [p.text for p in doc.paragraphs if p.style.name == "Excerpt Heading"]
    assert len(heads) == len(res.excerpts) and all("PDF page 1" in h for h in heads)
    info = [p.text for p in doc.paragraphs if "highlights found" in p.text][0]
    assert f"{res.found} highlights found" in info


def test_combined_file_structure(tmp_path):
    a = make(tmp_path / "a.pdf")
    b = make(tmp_path / "b.pdf", (("lack the democratic standing", YELLOW),), text=PARA2)
    out = tmp_path / "c.docx"
    write_combined(out, [("a.pdf", extract(a)), ("b.pdf", extract(b))])
    doc = Document(out)
    heads = [p for p in doc.paragraphs if p.style.name == "Heading 1"]
    assert [h.text for h in heads] == ["a.pdf", "b.pdf"]
    assert heads[1].paragraph_format.page_break_before is True


def test_title_only_when_real(tmp_path):
    for meta, expected in (("Microsoft Word - draft.docx", ""), ("A Real Article Title", "A Real Article Title")):
        d = new_doc()
        d.new_page()
        d.set_metadata({"title": meta})
        put(d[0], 72, 100, PARA1)
        hl(d[0], "written limits")
        assert extract(save(d, tmp_path / f"{len(expected)}.pdf")).title == expected


# ---------------------------------------------------------------- Layer 3
def sha(path):
    with open(path, "rb") as f:
        return hashlib.sha256(f.read()).hexdigest()


def test_inputs_unchanged_and_rerun_never_overwrites(tmp_path):
    src = tmp_path / "in"
    src.mkdir()
    pdf = make(src / "a.pdf")
    before = sha(pdf)
    out = tmp_path / "out"
    rows = cli.run([str(src)], str(out), settings(), log=quiet)
    first = rows[0]["output"]
    assert os.path.exists(first) and sha(pdf) == before
    doc = Document(first)  # the user edits the Word file
    doc.add_paragraph("MY MANUAL EDIT")
    doc.save(first)
    edited = sha(first)
    rows2 = cli.run([str(src)], str(out), settings(), log=quiet)
    assert sha(first) == edited and sha(pdf) == before
    assert rows2[0]["output"] != first and os.path.exists(rows2[0]["output"])
    assert "left untouched" in rows2[0]["detail"]


def test_duplicate_source_names_get_distinct_outputs(tmp_path):
    for sub in ("one", "two"):
        (tmp_path / "in" / sub).mkdir(parents=True)
        make(tmp_path / "in" / sub / "Notes.pdf")
    rows = cli.run([str(tmp_path / "in")], str(tmp_path / "out"), settings(), log=quiet)
    outs = sorted(os.path.basename(r["output"]) for r in rows)
    assert len(set(outs)) == 2 and all(o.endswith("(highlights).docx") for o in outs)
    assert any(o.startswith("one - ") for o in outs) and any(o.startswith("two - ") for o in outs)


def test_mixed_folder_each_problem_reported_and_good_file_exported(tmp_path):
    src = tmp_path / "in"
    src.mkdir()
    make(src / "good.pdf")
    (src / "corrupt.pdf").write_bytes(b"this is not a pdf at all")
    (src / "empty.pdf").write_bytes(b"")
    d = new_doc()
    d.new_page()
    put(d[0], 72, 100, PARA1)
    hl(d[0], "written limits")
    d.save(str(src / "locked.pdf"), encryption=pymupdf.PDF_ENCRYPT_AES_256, user_pw="x", owner_pw="y")
    d.close()
    d = new_doc()  # image-only scan: a picture and a highlight, no text
    p = d.new_page()
    pix = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 200, 100), False)
    pix.set_rect(pix.irect, (200, 200, 200))
    p.insert_image(pymupdf.Rect(50, 50, 400, 300), pixmap=pix)
    p.add_highlight_annot(pymupdf.Rect(60, 60, 200, 80))
    save(d, src / "scan.pdf")
    d = new_doc()
    d.new_page()
    put(d[0], 72, 100, PARA1)
    save(d, src / "plain.pdf")
    d = new_doc()  # "flattened": a yellow rectangle painted behind text, no annotation
    p = d.new_page()
    p.draw_rect(pymupdf.Rect(70, 88, 300, 104), color=None, fill=(1, 1, 0))
    put(p, 72, 100, PARA1)
    save(d, src / "flat.pdf")
    rows = {os.path.basename(r["file"]): r for r in
            cli.run([str(src)], str(tmp_path / "out"), settings(), log=quiet)}
    assert rows["good.pdf"]["status"] == "exported" and os.path.exists(rows["good.pdf"]["output"])
    assert rows["corrupt.pdf"]["status"] == "unreadable"
    assert rows["empty.pdf"]["status"] == "unreadable" and "empty" in rows["empty.pdf"]["detail"]
    assert rows["locked.pdf"]["status"] == "locked"
    assert rows["scan.pdf"]["status"] == "no text layer"
    assert rows["plain.pdf"]["status"] == "no highlights" and "painted" not in rows["plain.pdf"]["detail"]
    assert rows["flat.pdf"]["status"] == "no highlights" and "painted" in rows["flat.pdf"]["detail"]
    assert len([f for f in os.listdir(tmp_path / "out" / "Reports") if f.endswith(".txt")]) == 1


def test_report_counts_add_up(tmp_path):
    d = new_doc()
    p = d.new_page()
    put(p, 72, 100, PARA1)
    hl(p, "written limits")
    hl(p, "framers")
    p.add_highlight_annot(pymupdf.Rect(300, 600, 400, 614))
    pdf = save(d, tmp_path / "a.pdf")
    r = cli.run([pdf], str(tmp_path / "out"), settings(), log=quiet)[0]
    assert r["found"] == 3 and r["found"] == r["exported"] + r["no_text"]


def test_combined_option_creates_combined_file(tmp_path):
    (tmp_path / "in").mkdir()
    for n in "ab":
        make(tmp_path / "in" / f"{n}.pdf")
    cli.run([str(tmp_path / "in")], str(tmp_path / "out"), settings(combined=True), log=quiet)
    names = os.listdir(tmp_path / "out")
    assert "_Combined highlights.docx" in names
    assert sum(n.endswith("(highlights).docx") and not n.startswith("_") for n in names) == 2


def test_no_network_access(tmp_path, monkeypatch):
    def boom(*a, **k):
        raise AssertionError("network access attempted")
    monkeypatch.setattr(socket.socket, "connect", boom)
    monkeypatch.setattr(socket, "create_connection", boom)
    pdf = make(tmp_path / "a.pdf")
    assert cli.run([pdf], str(tmp_path / "out"), settings(), log=quiet)[0]["status"] == "exported"


def test_unreadable_file_message_is_plain(tmp_path):
    with pytest.raises(PdfProblem) as e:
        extract(str(tmp_path / "does-not-exist.pdf"))
    assert e.value.status == "cannot open file" and "Always keep on this device" in e.value.detail


# ---------------------------------------------------------------- privacy rule
def test_repo_contains_no_personal_references():
    """The shipped tool must not mention the owner's account, paths, or any sample document."""
    root = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    out = subprocess.run(["git", "ls-files"], cwd=root, capture_output=True, text=True)
    files = [f for f in out.stdout.split("\n") if f] if out.returncode == 0 else []
    files = files or [os.path.relpath(os.path.join(d, f), root) for d, _, fs in os.walk(root) for f in fs
                      if not any(x in d for x in (".git", "test-samples", "__pycache__", ".pytest_cache"))]
    banned = [r"C:\\Users", r"C:/Users", r"OneDrive\\Desktop", re.escape(getpass.getuser()),
              re.escape(os.path.basename(os.path.expanduser("~")))]
    sample_dir = os.path.join(root, "test-samples")
    if os.path.isdir(sample_dir):
        for dp, _, fs in os.walk(sample_dir):
            for f in fs:
                stem, ext = os.path.splitext(f)
                if ext.lower() == ".pdf" and len(stem) >= 8:
                    banned.append(re.escape(stem))
    this = os.path.basename(__file__)
    bad = []
    for rel in files:
        if os.path.basename(rel) in (this, ".git") or rel.startswith("test-samples"):
            continue
        try:
            with open(os.path.join(root, rel), encoding="utf-8") as f:
                text = f.read()
        except (OSError, UnicodeDecodeError):
            continue
        bad += [(rel, pat) for pat in banned if re.search(pat, text, re.I)]
    assert not bad, bad


def test_very_long_names_and_folders_still_save(tmp_path):
    """A long PDF name inside a deep folder must not exceed the Windows path limit."""
    long_name = "01 - " + "DELICATE TASK CONTENT MODERATION AND INTERMEDIARY LIABILITY IN A POST-DSA " * 1
    src = tmp_path / "in"
    src.mkdir()
    make(src / (long_name.strip() + ".pdf"))
    out = tmp_path
    for part in ("a" * 40, "b" * 40, "c" * 40):
        out = out / part
    deep = str(out)
    rows = cli.run([str(src)], deep, settings(), log=quiet)
    assert rows[0]["status"] == "exported", rows[0]
    assert os.path.exists(rows[0]["output"]) and len(rows[0]["output"]) < 260


def test_failed_file_reports_no_exported_counts(tmp_path, monkeypatch):
    pdf = make(tmp_path / "a.pdf")
    monkeypatch.setattr(cli.write_docx, "write_single", lambda *a, **k: (_ for _ in ()).throw(OSError("disk full")))
    r = cli.run([pdf], str(tmp_path / "out"), settings(), log=quiet)[0]
    assert r["status"] == "failed" and r["exported"] == 0 and r["found"] == 0


def test_word_files_folder_holds_only_word_files(tmp_path):
    (tmp_path / "in").mkdir()
    for n in "ab":
        make(tmp_path / "in" / f"{n}.pdf")
    cli.run([str(tmp_path / "in")], str(tmp_path / "out"), settings(combined=True), log=quiet)
    top = os.listdir(tmp_path / "out")
    assert sorted(f for f in top if not os.path.isdir(tmp_path / "out" / f)) == \
        sorted(f for f in top if f.endswith(".docx")) and "Reports" in top
    reports = os.listdir(tmp_path / "out" / "Reports")
    assert any(f.endswith(".txt") for f in reports) and any(f.endswith(".csv") for f in reports)
